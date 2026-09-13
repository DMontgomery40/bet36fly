# ruff: noqa: E402
"""Independent saved-array/workbook analysis. Never imports or runs the model."""

import time

_START = time.monotonic()

import argparse
import hashlib
import json
import os
from pathlib import Path
import signal

import numpy as np

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "incentive-circuit-source-2026-09-13"
WORKBOOK = SOURCE / "src/incentive/data/handler2019/Handler2019_Fig2Fig5_Data.xlsx"
WORKBOOK_SHA = "4db9b76cd0c41dc014c8d00a9e359bbd4165a7984095cc35d1932f9ee7f501dc"
ISIS = [-6., -1.2, -.6, 0., .5, 6.]
SOURCE_HASHES = {
    "src/incentive/handler.py": "7a943ec47bfc0db20ab2679f8b5f70f1a39dd9d646dbdac33fd99b4aecf0912a",
    "src/incentive/models_base.py": "0113fedc0f15eb8c1f5c0f94a4bd566e1723a0aac34e4ab2ca8538363187f0a7",
    "src/incentive/circuit.py": "94139014ada70594545becc730fb4e9c86e0aab584d39cf993e54d8e37acbb9a",
    "examples/run_handler_2019.py": "cb4d4109083b162ee1547f385409955d65d1bedda00de4ecdc82a038cde75e53",
    "README.md": "7bfb94640ce09dde570f08a48dd44f1a7c14dd49a8dbc64745c6bf5678580062",
    "LICENSE": "3972dc9744f6499f0f9b2dbf76696f2ae7ad8af9b23dde66d6af86c9dfb36986",
}
BASE_FIELDS = ("time", "cs", "us", "k", "d1", "d2", "m", "w", "dR1", "dR2",
               "w_before", "delta_w", "internal_w")
ARRAY_FIELDS = BASE_FIELDS + tuple(prefix + k for prefix in ("independent_", "difference_")
                                   for k in BASE_FIELDS if k != "time")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def window(t, y, lo, hi, *, closed):
    t, y = np.asarray(t, float), np.asarray(y, float)
    if (t.ndim != 1 or not t.size or y.shape != t.shape or
            not np.isfinite(t).all() or not np.isfinite(y).all() or
            np.any(np.diff(t) <= 0) or not np.isfinite([lo, hi]).all() or hi <= lo):
        raise ValueError("Invalid window data or domain")
    selected = (t >= lo) & ((t <= hi) if closed else (t < hi))
    n = int(selected.sum())
    return {"requested": [float(lo), float(hi)], "closed": closed, "samples": n,
            "complete": bool(t[0] <= lo and t[-1] >= hi),
            "mean": float(np.mean(y[selected])) if n else None,
            "first_selected": float(t[selected][0]) if n else None,
            "last_selected": float(t[selected][-1]) if n else None}


def normalize(values):
    x = np.asarray(values, float)
    if (x.ndim != 2 or min(x.shape) < 1 or x.shape[0] < 2 or
            not np.isfinite(x).all() or np.any(np.ptp(x, axis=0) <= 0)):
        raise ValueError("Invalid or degenerate normalization")
    with np.errstate(over="raise", invalid="raise", divide="raise"):
        z = (x - x.min(axis=0)) / np.ptp(x, axis=0)
    return {"mean": z.mean(axis=1).tolist(),
            "sem": (z.std(axis=1, ddof=1) / np.sqrt(x.shape[1])).tolist() if x.shape[1] > 1 else None,
            "per_preparation": z.tolist(), "minimum": x.min(axis=0).tolist(),
            "range": np.ptp(x, axis=0).tolist()}


def pearson(x, y):
    x, y = np.asarray(x, float), np.asarray(y, float)
    if x.ndim != 1 or x.shape != y.shape or x.size < 2 or not np.isfinite([x, y]).all():
        raise ValueError("Invalid correlation inputs")
    if np.ptp(x) == 0 or np.ptp(y) == 0:
        return None
    return float(np.corrcoef(x, y)[0, 1])


def analyze_cases(cases):
    if len(cases) != 6 or sorted(c["isi"] for c in cases) != ISIS:
        raise ValueError("Six unique canonical conditions required")
    rows = []
    for c in sorted(cases, key=lambda c: c["isi"]):
        isi, t = c["isi"], c["time"]
        row = {"isi": isi}
        for kind, lo, hi, closed, field, sign in [
            ("literal_ER", -7, 1, False, "dR1", -1),
            ("literal_cAMP", isi, isi + 4, False, "dR2", 1),
            ("primary_ER", 0, 1, True, "dR1", -1),
            ("primary_cAMP", isi, isi + 4, True, "dR2", 1),
        ]:
            row[kind] = window(t, sign * np.asarray(c[field]), lo, hi, closed=closed)
        row.update(final_internal_w=float(c["internal_w"][-1]),
                   final_recorded_w=float(c["w"][-1]),
                   internal_weight_change=float(c["internal_w"][-1] - 1),
                   signed_increment_sum=float(np.sum(c["delta_w"])),
                   reporter_proxy_sum=float(np.sum(-c["dR1"] - c["dR2"])))
        rows.append(row)
    er = normalize([[r["literal_ER"]["mean"]] for r in rows])
    ca = normalize([[r["literal_cAMP"]["mean"]] for r in rows])
    complete = [i for i, r in enumerate(rows) if r["primary_ER"]["complete"] and r["primary_cAMP"]["complete"]]
    return {"rows": rows, "literal_model": {"ER": er, "cAMP": ca,
                "normalized_contrast": (np.array(er["mean"]) - ca["mean"]).tolist()},
            "primary_model": {"complete_condition_indices": complete,
                "normalized_contrast": None, "correlation": None,
                "reason": "The +6s cAMP window requires10s; saved source ends8s. No six-condition normalization or five-condition renormalization."}}


def read_saved(run, expected_identity_sha):
    run = Path(run)
    if (run / "terminal-error.json").exists() or sha(run / "identity.json") != expected_identity_sha:
        raise ValueError("Failed terminal or identity hash mismatch")
    status = json.loads((run / "status.json").read_text())
    identity = json.loads((run / "identity.json").read_text())
    if (status.get("status") != "completed" or status.get("completed_cases") != 6 or
        status.get("source_unchanged") is not True or not 0 <= status.get("wall_seconds", -1) < 30 or
        len(status.get("rows", [])) != 6):
        raise ValueError("Incomplete or out-of-cap source calculation")
    for key, expected in {"source": SOURCE_HASHES, "ISIS": ISIS, "taus": [100/3, 60., 104.],
                          "samples": 1001, "grid": [-7, 8], "cap_seconds": 30., "scope": "authored_source_only"}.items():
        if identity.get(key) != expected:
            raise ValueError("Wrong source protocol identity")
    hashes = {str(run / n): sha(run / n) for n in ("identity.json", "status.json", "summary.json")}
    expected = {str(SOURCE / n): h for n, h in SOURCE_HASHES.items()}
    expected.update(identity["bindings"])
    expected[str(run / "summary.json")] = status["summary_sha256"]
    cases = []
    for i, (isi, row) in enumerate(zip(ISIS, status["rows"])):
        name = f"case_{i:02d}.npz"
        if (row.get("index") != i or row.get("isi") != isi or row.get("file") != name or
                row.get("passed") is not True):
            raise ValueError("Wrong or failed case identity")
        expected[str(run / name)] = row["sha256"]
    for p, h in expected.items():
        hashes[p] = sha(p)
        if hashes[p] != h:
            raise ValueError("Input hash mismatch: " + p)
    for i, isi in enumerate(ISIS):
        with np.load(run / f"case_{i:02d}.npz", allow_pickle=False) as z:
            if set(z.files) != set(ARRAY_FIELDS):
                raise ValueError("Wrong array schema")
            c = {k: z[k].copy() for k in z.files}
        if any(x.shape != (1001,) or x.dtype != np.float64 or not np.isfinite(x).all() for x in c.values()):
            raise ValueError("Malformed or nonfinite saved arrays")
        t = np.linspace(-7., 8., 1001)
        if not (np.array_equal(c["time"], t) and
                np.array_equal(c["cs"], (t >= 0) & (t < .5)) and
                np.array_equal(c["us"], (t >= isi) & (t < isi + .6))):
            raise ValueError("Wrong recorded source clock or stimulus")
        c["isi"] = isi
        cases.append(c)
    return cases, json.loads((run / "summary.json").read_text()), hashes


def read_workbook(path=WORKBOOK):
    if sha(path) != WORKBOOK_SHA:
        raise ValueError("Workbook hash mismatch")
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True, read_only=False)
    means, literal, counts = {}, {}, {}
    for branch, sheet, starts, nprep, ntime, offset in [
        ("cAMP", "Fig5D_cAMP timecourse", [3, 11, 19, 27, 35, 43], 6, 300, 10),
        ("ER", "Fig5D_ERGCaMP timecourse", [3, 12, 21, 30, 39, 48], 7, 200, 6),
    ]:
        ws = book[sheet]
        means[branch], literal[branch], counts[branch] = [], [], []
        for isi, start in zip(ISIS, starts):
            t = np.array([ws.cell(r, start - 1).value for r in range(4, 4 + ntime)], float)
            data = np.array([[ws.cell(r, col).value for col in range(start, start + nprep)]
                             for r in range(4, 4 + ntime)], float)
            lo, hi = (offset + isi, offset + isi + 4) if branch == "cAMP" else (6, 7)
            rows = [window(t, data[:, j], lo, hi, closed=True) for j in range(nprep)]
            v = [r["mean"] for r in rows]
            if any(not r["complete"] for r in rows) or not np.allclose(v, [ws.cell(2, col).value for col in range(start, start + nprep)], rtol=0, atol=1e-14):
                raise ValueError("Primary workbook means or coverage differ")
            means[branch].append(v)
            literal[branch].append(v if branch == "cAMP" else [window(t, data[:, j], -1, 7, closed=True)["mean"] for j in range(nprep)])
            counts[branch].append(rows[0]["samples"])
    ws = book["Fig5E_(-)ERGCaMPnorm - cAMPnorm"]
    rounded = {"cAMP": [[ws.cell(r, c).value for c in range(2, 8)] for r in range(3, 9)],
               "ER": [[ws.cell(r, c).value for c in range(2, 9)] for r in range(21, 27)]}

    def branches(m):
        er, ca = normalize(-np.array(m["ER"])), normalize(m["cAMP"])
        return {"ER": er, "cAMP": ca,
                "contrast": (np.array(er["mean"]) - ca["mean"]).tolist(),
                "contrast_sem": np.hypot(er["sem"], ca["sem"]).tolist()}

    rounded_result = branches(rounded)
    if not np.allclose(rounded_result["contrast"], [ws.cell(r, 14).value for r in range(39, 45)], rtol=0, atol=1e-12):
        raise ValueError("Rounded workbook contrast differs")
    mb = book["Fig2D-2E_gamma4MBON"]
    mbon = []
    for i, isi in enumerate(ISIS):
        r = i + 3
        if mb.cell(r, 80).value != isi:  # CB contains numeric ISI.
            raise ValueError("MBON condition order differs")
        values = [mb.cell(r, c).value for c in range(81, 87) if mb.cell(r, c).value is not None]
        if len(values) != [5, 5, 6, 5, 5, 5][i] or not np.isfinite(values).all():
            raise ValueError("MBON preparation count or data differ")
        mbon.append({"isi": isi, "preparations": len(values), "mean": float(np.mean(values)),
                     "sem": float(np.std(values, ddof=1) / np.sqrt(len(values))), "values": values})
    book.close()
    return {"primary_raw_means": means, "primary_samples": counts,
            "primary_full_precision": branches(means), "primary_rounded_sheet": rounded_result,
            "literal_driver_experimental": branches(literal), "direct_MBON": mbon,
            "literal_ER_available_samples": 71, "literal_ER_requested_raw_window": [-1, 7]}


def compare_summary(result, saved):
    rows, literal = result["rows"], result["literal_model"]
    expected = {"ER_means": [r["literal_ER"]["mean"] for r in rows],
        "cAMP_means": [r["literal_cAMP"]["mean"] for r in rows],
        "ER_window_samples": [r["literal_ER"]["samples"] for r in rows],
        "cAMP_window_samples": [r["literal_cAMP"]["samples"] for r in rows],
        "ER_normalized": literal["ER"]["mean"], "cAMP_normalized": literal["cAMP"]["mean"],
        "normalized_contrast": literal["normalized_contrast"],
        "final_internal_w": [r["final_internal_w"] for r in rows],
        "final_recorded_w": [r["final_recorded_w"] for r in rows]}
    if set(saved) != set(expected) or any(np.asarray(saved[k]).shape != (6,) or
            not np.allclose(saved[k], v, rtol=0, atol=1e-12) for k, v in expected.items()):
        raise ValueError("Independent saved-summary comparison failed")


def write_json(path, value):
    with Path(path).open("x") as f:
        json.dump(value, f, indent=2, allow_nan=False)
        f.write("\n")
        f.flush()
        os.fsync(f.fileno())


def execute(run, identity_sha, out, *, started=_START, clock=time.monotonic):
    out = Path(out)
    out.mkdir(exist_ok=False)

    def guard():
        elapsed = clock() - started
        if not 0 <= elapsed < 30:
            raise TimeoutError("Independent saved-analysis30s cap")
        return elapsed

    try:
        write_json(out / "started.json", {"identity_sha256": identity_sha, "elapsed_seconds": guard()})
        cases, saved, hashes = read_saved(run, identity_sha)
        for p in (WORKBOOK, Path(__file__), HERE / "test_analyze_incentive_handler_saved.py",
                  HERE / "incentive-handler-saved-analysis-preregistration-2026-09-13.md"):
            hashes[str(p)] = sha(p)
        guard()
        result = analyze_cases(cases)
        compare_summary(result, saved)
        result["workbook"] = data = read_workbook()
        model = result["literal_model"]["normalized_contrast"]
        result["descriptive_comparisons"] = {
            "literal_model_vs_literal_experimental_contrast_r": pearson(model, data["literal_driver_experimental"]["contrast"]),
            "literal_model_vs_primary_rounded_reporter_contrast_r": pearson(model, data["primary_rounded_sheet"]["contrast"]),
            "literal_model_vs_direct_MBON_mean_r": pearson(model, [r["mean"] for r in data["direct_MBON"]]),
            "primary_five_complete_conditions": [{"isi": ISIS[i],
                "model_ER_raw": result["rows"][i]["primary_ER"]["mean"],
                "model_cAMP_raw": result["rows"][i]["primary_cAMP"]["mean"],
                "data_ER_raw_mean": -float(np.mean(data["primary_raw_means"]["ER"][i])),
                "data_cAMP_raw_mean": float(np.mean(data["primary_raw_means"]["cAMP"][i]))}
                for i in result["primary_model"]["complete_condition_indices"]],
            "scope": "Descriptive, same data used by authors for timing selection; no acceptance threshold, calibrated units, causal receptor validation or MaleCNS qualification."}
        result["input_hashes"] = hashes
        guard()
        if any(sha(p) != h for p, h in hashes.items()) or (Path(run) / "terminal-error.json").exists():
            raise ValueError("Inputs changed during analysis")
        write_json(out / "analysis.json", result)
        write_json(out / "completion.json", {"status": "complete", "analysis_sha256": sha(out / "analysis.json"),
            "source_identity_sha256": identity_sha, "elapsed_seconds": guard(), "cases": 6,
            "model_executions": 0, "fits": 0})
        guard()
        return result
    except BaseException as error:
        write_json(out / "terminal-error.json", {"status": "failed", "error": type(error).__name__ + ": " + str(error),
                                                "elapsed_seconds": clock() - started})
        raise


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--run", type=Path, required=True)
    p.add_argument("--identity-sha", required=True)
    p.add_argument("--out", type=Path, required=True)
    args = p.parse_args()

    def expired(*_):
        raise TimeoutError("Independent saved-analysis30s wall cap")

    signal.signal(signal.SIGALRM, expired)
    remaining = 30 - (time.monotonic() - _START)
    if remaining <= 0:
        raise TimeoutError("Imports exhausted audit cap")
    signal.setitimer(signal.ITIMER_REAL, remaining)
    try:
        execute(args.run, args.identity_sha, args.out)
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)


if __name__ == "__main__":
    main()
